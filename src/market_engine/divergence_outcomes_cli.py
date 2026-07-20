from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl.chart import LineChart, Reference

from market_engine.cli import _excel_safe
from market_engine.config import PROJECT_ROOT
from market_engine.evaluation.divergence_outcomes import run_divergence_outcome_laboratory


def _resolve(path_text: str) -> Path:
    path = Path(path_text)
    return path if path.is_absolute() else PROJECT_ROOT / path


def _read(path: Path) -> pd.DataFrame:
    if path.suffix.lower() == ".csv":
        return pd.read_csv(path)
    return pd.read_excel(path)


def _parse_horizons(value: str) -> tuple[int, ...]:
    try:
        horizons = tuple(int(item.strip()) for item in value.split(",") if item.strip())
    except ValueError as exc:
        raise argparse.ArgumentTypeError("horizons debe ser una lista de enteros") from exc
    if not horizons or any(item < 1 for item in horizons):
        raise argparse.ArgumentTypeError("horizons debe contener enteros positivos")
    return horizons


def _add_chart(
    writer: pd.ExcelWriter,
    sheet_name: str,
    title: str,
    columns: list[int],
    anchor: str,
) -> None:
    sheet = writer.book[sheet_name]
    chart = LineChart()
    chart.title = title
    chart.y_axis.title = title
    chart.x_axis.title = "Fecha"
    chart.height = 8
    chart.width = 16
    categories = Reference(sheet, min_col=1, min_row=2, max_row=sheet.max_row)
    for column in columns:
        data = Reference(sheet, min_col=column, min_row=1, max_row=sheet.max_row)
        chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    sheet.add_chart(chart, anchor)


def _evaluate(args: argparse.Namespace) -> int:
    sample_path = _resolve(args.sample)
    reference_path = _resolve(args.reference)
    if not sample_path.exists():
        raise FileNotFoundError(f"No existe la muestra: {sample_path}")
    if not reference_path.exists():
        raise FileNotFoundError(f"No existe la referencia: {reference_path}")

    result = run_divergence_outcome_laboratory(
        _read(sample_path),
        _read(reference_path),
        horizons=args.horizons,
        smoothing_window=args.smoothing_window,
        spread_window=args.spread_window,
        extreme_zscore=args.extreme_zscore,
        minimum_episode_observations=args.minimum_episode_observations,
    )

    print("\nDIVERGENCE OUTCOME LABORATORY v1.0.0a5")
    print(result.summary.to_string(index=False))

    output = _resolve(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        _excel_safe(result.summary).to_excel(writer, sheet_name="Resumen", index=False)
        _excel_safe(result.state_outcomes).to_excel(writer, sheet_name="Resultados_Estados", index=False)
        _excel_safe(result.transition_matrix).to_excel(writer, sheet_name="Transiciones", index=False)
        _excel_safe(result.episode_outcomes).to_excel(writer, sheet_name="Resultados_Episodios", index=False)
        _excel_safe(result.daily_outcomes).to_excel(writer, sheet_name="Resultados_Diarios", index=False)
        _excel_safe(result.divergence_result.state_summary).to_excel(
            writer, sheet_name="Estados_Base", index=False
        )

        sheet = writer.book["Resultados_Diarios"]
        headers = {cell.value: cell.column for cell in sheet[1]}
        _add_chart(
            writer,
            "Resultados_Diarios",
            "Geometría normalizada",
            [headers["sample_market_cap_index"], headers["reference_market_cap_index"]],
            "AZ2",
        )
        _add_chart(
            writer,
            "Resultados_Diarios",
            "Leadership spread",
            [headers["leadership_spread"], headers["rolling_spread_mean"]],
            "AZ20",
        )
        horizon = max(args.horizons)
        _add_chart(
            writer,
            "Resultados_Diarios",
            f"Retornos futuros a {horizon} sesiones",
            [
                headers[f"sample_return_forward_{horizon}"],
                headers[f"reference_return_forward_{horizon}"],
            ],
            "AZ38",
        )
        _add_chart(
            writer,
            "Resultados_Diarios",
            f"Cambio futuro del spread a {horizon} sesiones",
            [headers[f"spread_change_forward_{horizon}"]],
            "AZ56",
        )

    print(f"Excel generado: {output}")
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(prog="market-divergence-outcomes")
    parser.add_argument("--sample", required=True)
    parser.add_argument("--reference", required=True)
    parser.add_argument("--horizons", type=_parse_horizons, default=(1, 5, 10, 20))
    parser.add_argument("--smoothing-window", type=int, default=5)
    parser.add_argument("--spread-window", type=int, default=20)
    parser.add_argument("--extreme-zscore", type=float, default=2.0)
    parser.add_argument("--minimum-episode-observations", type=int, default=3)
    parser.add_argument("--output", default="divergence_outcomes_v100a5.xlsx")
    parser.set_defaults(handler=_evaluate)
    return parser


def main() -> int:
    args = build_parser().parse_args()
    return int(args.handler(args))


if __name__ == "__main__":
    raise SystemExit(main())
