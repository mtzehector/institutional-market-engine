from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl.chart import LineChart, Reference

from market_engine.cli import _excel_safe
from market_engine.config import PROJECT_ROOT
from market_engine.evaluation.sentinel_divergence import run_sentinel_divergence_laboratory


def _resolve(path_text: str) -> Path:
    path = Path(path_text)
    return path if path.is_absolute() else PROJECT_ROOT / path


def _read(path: Path) -> pd.DataFrame:
    if path.suffix.lower() == ".csv":
        return pd.read_csv(path)
    return pd.read_excel(path)


def _add_chart(writer: pd.ExcelWriter, sheet_name: str, title: str, columns: list[int], anchor: str) -> None:
    sheet = writer.book[sheet_name]
    chart = LineChart()
    chart.title = title
    chart.y_axis.title = title
    chart.x_axis.title = "Fecha"
    chart.height = 8
    chart.width = 16
    categories = Reference(sheet, min_col=1, min_row=2, max_row=sheet.max_row)
    for column in columns:
        data = Reference(sheet, min_col=column, min_row=1, max_row=sheet.max_row)
        chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    sheet.add_chart(chart, anchor)


def _evaluate(args: argparse.Namespace) -> int:
    sample_path = _resolve(args.sample)
    reference_path = _resolve(args.reference)
    if not sample_path.exists():
        raise FileNotFoundError(f"No existe la muestra: {sample_path}")
    if not reference_path.exists():
        raise FileNotFoundError(f"No existe la referencia: {reference_path}")

    result = run_sentinel_divergence_laboratory(
        _read(sample_path),
        _read(reference_path),
        smoothing_window=args.smoothing_window,
        spread_window=args.spread_window,
        extreme_zscore=args.extreme_zscore,
        minimum_episode_observations=args.minimum_episode_observations,
    )

    print("\nSENTINEL LEADERSHIP AND DIVERGENCE LABORATORY v1.0.0a4")
    print(result.summary.to_string(index=False))

    output = _resolve(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        _excel_safe(result.summary).to_excel(writer, sheet_name="Resumen", index=False)
        _excel_safe(result.daily).to_excel(writer, sheet_name="Divergencia_Diaria", index=False)
        _excel_safe(result.episodes).to_excel(writer, sheet_name="Episodios_Divergencia", index=False)
        _excel_safe(result.state_summary).to_excel(writer, sheet_name="Estados", index=False)
        _excel_safe(result.extremes).to_excel(writer, sheet_name="Extremos", index=False)

        sheet = writer.book["Divergencia_Diaria"]
        headers = {cell.value: cell.column for cell in sheet[1]}
        _add_chart(writer, "Divergencia_Diaria", "Geometría normalizada", [headers["sample_market_cap_index"], headers["reference_market_cap_index"]], "AA2")
        _add_chart(writer, "Divergencia_Diaria", "Leadership spread", [headers["leadership_spread"], headers["rolling_spread_mean"]], "AA20")
        _add_chart(writer, "Divergencia_Diaria", "Velocidad y aceleración", [headers["spread_velocity"], headers["spread_acceleration"]], "AA38")
        _add_chart(writer, "Divergencia_Diaria", "Spread z-score", [headers["spread_zscore"]], "AA56")
        _add_chart(writer, "Divergencia_Diaria", "Drawdowns comparados", [headers["sample_drawdown_pct"], headers["reference_drawdown_pct"]], "AA74")

    print(f"Excel generado: {output}")
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(prog="market-sentinel-divergence")
    parser.add_argument("--sample", required=True)
    parser.add_argument("--reference", required=True)
    parser.add_argument("--sample-label", default="SAMPLE")
    parser.add_argument("--reference-label", default="REFERENCE")
    parser.add_argument("--smoothing-window", type=int, default=5)
    parser.add_argument("--spread-window", type=int, default=20)
    parser.add_argument("--extreme-zscore", type=float, default=2.0)
    parser.add_argument("--minimum-episode-observations", type=int, default=3)
    parser.add_argument("--output", default="sentinel_divergence_v100a4.xlsx")
    parser.set_defaults(handler=_evaluate)
    return parser


def main() -> int:
    args = build_parser().parse_args()
    return int(args.handler(args))


if __name__ == "__main__":
    raise SystemExit(main())
